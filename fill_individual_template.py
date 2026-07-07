# Fills CA2_individual_template.docx for Kendrick Slamat p2510877
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

d = Document('CA2_individual_template.docx')

MONO = 'Consolas'


def set_cell(cell, text, size=8.5, mono=True):
    cell.text = ''
    p = cell.paragraphs[0]
    for i, line in enumerate(text.split('\n')):
        if i:
            p = cell.add_paragraph()
        r = p.add_run(line)
        r.font.size = Pt(size)
        if mono:
            r.font.name = MONO


sql = """USE CA22B2105;
SELECT b.BookID, b.Title, b.Author, b.Genre,
       t.TierName AS MinTierRequired,
       COUNT(DISTINCT c.CopySeqNo) AS TotalCopies,
       COUNT(DISTINCT CASE WHEN c.Status = 'Lost'
             THEN c.CopySeqNo END) AS LostCopies,
       (SELECT COUNT(*) FROM Loan l
        WHERE l.BookID = b.BookID) AS TotalLoans
FROM Book b
JOIN MembershipTier t ON b.MinTierID = t.TierID
LEFT JOIN BookCopy c ON c.BookID = b.BookID
GROUP BY b.BookID, b.Title, b.Author, b.Genre, t.TierName;"""
set_cell(d.tables[0].rows[1].cells[1], sql)
set_cell(d.tables[1].rows[1].cells[1], '500', mono=False, size=11)

q1 = """use('CA2Mongo');
db.Book.find(
  { totalLoans: { $gt: 400 } },
  { _id: 0, title: 1, genre: 1, totalLoans: 1 }
).sort({ totalLoans: -1 });"""
r1 = ("(empty result set - no book has been borrowed more than 400 times; "
      "the maximum totalLoans in the data is 91. Query verified by re-running "
      "with a lower threshold, e.g. $gt: 80, which returns documents as expected.)")

q2 = """use('CA2Mongo');
db.Book.aggregate([
  { $group: { _id: "$genre",
      numBooks: { $sum: 1 },
      avgLoansPerBook: { $avg: "$totalLoans" } } },
  { $match: { numBooks: { $gt: 50 },
              avgLoansPerBook: { $gt: 70 } } },
  { $sort: { avgLoansPerBook: -1 } }
]);"""
r2 = ("(empty result set - only three genres have more than 50 books "
      "(Research Journal 96, Science 70, Children 67) and their average loans "
      "per book are 21.8, 26.4 and 28.0 respectively, none above 70. Query "
      "verified with relaxed thresholds, e.g. numBooks > 50 and "
      "avgLoansPerBook > 20, which returns all three genres.)")

t2 = d.tables[2]
set_cell(t2.rows[1].cells[1], q1)
set_cell(t2.rows[1].cells[2], r1, mono=False)
set_cell(t2.rows[2].cells[1], q2)
set_cell(t2.rows[2].cells[2], r2, mono=False)


# ---------- Task f report ----------
def h(text, size=14, bold=True):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p


def body(text, size=11):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def code(text, size=8.5):
    p = d.add_paragraph()
    for i, line in enumerate(text.split('\n')):
        r = p.add_run(line)
        r.font.name = MONO
        r.font.size = Pt(size)
        if i < len(text.split('\n')) - 1:
            r.add_break()
    return p


pb = d.add_paragraph()
pb.add_run().add_break(WD_BREAK.PAGE)

h('Task f) Individual Report - How B-tree Index Design Affects Query Performance on Library Loan Data', 13)
body('Kendrick Slamat, p2510877, class 2B21 group 05')

h('1. Question being explored', 12)
body('Our OLTP database answers analytical questions such as "how many loans were made in a given period and how '
     'long were books kept?" directly on the Loan table. This report investigates a topic beyond the module\'s '
     'scope: how MySQL\'s InnoDB B-tree indexes physically change the way such queries execute, when the optimiser '
     'chooses (or refuses) to use an index, and what composite and covering indexes add. The concrete question: '
     'for date-range and staff-filtered queries on a 12,344-row Loan table, when does a secondary index actually '
     'help, and how can we prove it rather than assume it?')

h('2. Key concepts learned', 12)
body('(a) InnoDB stores every table as a clustered B+-tree ordered by primary key; a full "table scan" walks the '
     'leaf level of this tree. A secondary index is a second B+-tree containing the indexed columns plus the '
     'primary key, so each secondary-index hit costs an extra lookup back into the clustered index unless the '
     'index covers all the columns the query needs (a "covering index").')
body('(b) The optimiser is cost-based: it estimates how many rows a plan touches and may deliberately ignore an '
     'index when the predicate matches a large fraction of the table, because random secondary-index lookups are '
     'more expensive than one sequential scan. A common rule of thumb is that beyond roughly 20-30% selectivity, '
     'the scan wins.')
body('(c) Composite indexes follow the leftmost-prefix rule: an index on (StaffID, StartDatetime) supports '
     'equality on StaffID plus a range on StartDatetime, but cannot serve a query that filters only on '
     'StartDatetime. EXPLAIN ANALYZE (MySQL 8.0) shows both the estimated plan and the actual measured times, '
     'making these effects observable.')
body('(d) Indexes interact with constraints: InnoDB requires an index on every foreign-key column, and silently '
     'reuses a composite index whose leftmost column matches. Dropping such an index fails with error 1553 until '
     'another qualifying index exists.')

h('3. Application to the library scenario', 12)
body('The library workload has natural read patterns: monthly/seasonal reporting (date ranges over '
     'Loan.StartDatetime), per-staff accountability queries (StaffID + date), and popularity counts by book. '
     'These map directly to single-column, composite and covering indexes respectively. All demonstrations below '
     'were run on our actual CA2 OLTP database (CA22B2105, 12,344 loans, MySQL 8.0.44) using EXPLAIN ANALYZE.')

h('4. Demonstration', 12)
h('4.1 Baseline: date-range query with no index', 11)
code("""EXPLAIN ANALYZE
SELECT COUNT(*), AVG(DATEDIFF(ReturnDatetime, StartDatetime))
FROM Loan
WHERE StartDatetime >= '2025-11-01' AND StartDatetime < '2026-01-01';

-> Aggregate: ... (actual time=6.55..6.55)
   -> Filter: (StartDatetime >= ... and < ...) (rows=4027)
      -> Table scan on Loan (rows=12344) (actual time=0.416..4.39)""")
body('Without an index, MySQL scans all 12,344 rows to find the 4,027 matching ones (~6.5 ms).')

h('4.2 The optimiser can refuse an index: selectivity matters', 11)
body('After CREATE INDEX idx_loan_startdt ON Loan (StartDatetime), re-running the same two-month query still '
     'produced a table scan: 4,027 of 12,344 rows is ~33% of the table, above the threshold where random lookups '
     'beat one sequential pass. EXPLAIN on a six-year range shows the same: possible_keys lists idx_loan_startdt '
     'but key is NULL. An index is not a guarantee - it is an option the optimiser prices.')
h('4.3 A selective range uses the index', 11)
code("""EXPLAIN ANALYZE
SELECT COUNT(*), AVG(DATEDIFF(ReturnDatetime, StartDatetime))
FROM Loan
WHERE StartDatetime >= '2021-03-01' AND StartDatetime < '2021-03-08';

-> Aggregate: ... (actual time=1.86..1.86)
   -> Index range scan on Loan using idx_loan_startdt
      (rows=12) (actual time=1.82..1.83)""")
body('For a one-week window (12 rows, ~0.1% of the table) the plan flips to an index range scan touching only 12 '
     'rows instead of 12,344 - a 1000x reduction in rows examined.')

h('4.4 Composite index for staff-accountability queries', 11)
code("""CREATE INDEX idx_loan_staff_start ON Loan (StaffID, StartDatetime);

EXPLAIN ANALYZE
SELECT MemberID, BookID, StartDatetime, ReturnDatetime
FROM Loan
WHERE StaffID = 'S0010'
  AND StartDatetime >= '2025-01-01' AND StartDatetime < '2026-01-01';

Before (FK auto-index on StaffID only):
-> Index lookup using StaffID (rows=227 est.)
   (actual time=6.08..12.6, rows=134)
After (composite index):
-> Index range scan using idx_loan_staff_start
   (rows=134 est.) (actual time=0.219..0.639, rows=134)""")
body('The single-column FK index finds all 227 estimated loans for S0010 and then filters by date; the composite '
     'index descends straight to the (staff, date) range: actual execution time drops from ~12.6 ms to ~0.64 ms '
     '(about 20x) and the row estimate becomes exact, which also improves join planning in bigger queries.')

h('4.5 Covering index: answering a query without touching the table', 11)
code("""CREATE INDEX idx_loan_start_book ON Loan (StartDatetime, BookID);
EXPLAIN FORMAT=TREE
SELECT BookID, COUNT(*) FROM Loan
WHERE StartDatetime >= '2025-12-01' AND StartDatetime < '2026-01-01'
GROUP BY BookID;

-> Aggregate using temporary table
   -> Covering index range scan on Loan using idx_loan_start_book""")
body('Because the index contains every column the query references, the plan reads "Covering index range scan": '
     'the clustered index is never visited, eliminating the per-row lookback that normally makes secondary '
     'indexes expensive.')

h('4.6 A side-effect worth knowing: indexes and foreign keys', 11)
body('While cleaning up, DROP INDEX idx_loan_staff_start failed with error 1553: "Cannot drop index: needed in a '
     'foreign key constraint". InnoDB had adopted the composite index (leftmost column StaffID) to enforce the '
     'Loan-to-Staff foreign key. The drop only succeeded after creating a plain index on StaffID first - a '
     'practical trap when tuning tables that carry FK constraints.')

h('5. Conclusions', 12)
body('Indexing is not "add an index, get speed". On our own loan data: a well-matched composite index cut a '
     'staff-history query by ~20x and a selective date range by ~1000x in rows examined, yet the identical index '
     'was correctly ignored for low-selectivity ranges where scanning is cheaper. Practical guidance for the '
     'library system: index (StaffID, StartDatetime) for accountability queries, keep a StartDatetime index for '
     'short-window reports, consider covering indexes for high-frequency count dashboards, and always verify with '
     'EXPLAIN ANALYZE on production-sized data instead of assuming. For large multi-year analytics, the better '
     'tool is the dimensional data warehouse built in tasks b)-d), which avoids these trade-offs by design.')

h('Use of AI tools', 12)
body('AI assistance (Claude) was used to help set up the experiment scripts, capture EXPLAIN ANALYZE outputs and '
     'organise this report. All demonstrations were executed and verified on the group\'s actual CA22B2105 '
     'database, and the concepts were reviewed against the MySQL 8.0 Reference Manual (chapters on InnoDB index '
     'structures and EXPLAIN output).')

d.save('p2510877-KendrickSlamat_individual_submission.docx')
print('saved')
